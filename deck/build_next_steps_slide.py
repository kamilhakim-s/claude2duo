#!/usr/bin/env python3
"""Standalone builder for the single 'Next steps' slide (prompt experimentation +
codebase-agnostic evaluation). Renders one slide so it can be reviewed on its own and
later folded into the main deck.

Regenerate:  python3 deck/build_next_steps_slide.py
Output:      deck/next-steps-slide.pptx
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x5A, 0x6B, 0x7A)
ACCENT = RGBColor(0x0E, 0x63, 0xB0)
ACCENT_DARK = RGBColor(0x0A, 0x46, 0x7E)
ACCENT_LIGHT = RGBColor(0xE3, 0xEE, 0xF8)
GOOD = RGBColor(0x1E, 0x7A, 0x45)
WARN = RGBColor(0xB4, 0x5E, 0x12)
CARD_BG = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])


def set_text(tf, paras, align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def textbox(left, top, width, height):
    box = s.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box


def card(left, top, width, height, fill=CARD_BG, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def card_text(shp, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shp.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.12)
    set_text(tf, paras, align=align)


# header
rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.55), Inches(0.55), Inches(0.09))
rule.fill.solid()
rule.fill.fore_color.rgb = ACCENT
rule.line.fill.background()
rule.shadow.inherit = False
box = textbox(Inches(0.6), Inches(0.68), Inches(12.1), Inches(1.0))
set_text(box.text_frame, [
    [("NEXT STEPS", 12, True, ACCENT)],
    [("Experiment to find the right prompt — and prove each pack against its own codebase", 26, True, INK)],
])

# Why this is the next step (left) + The experiment (right)
c1 = card(Inches(0.6), Inches(1.85), Inches(6.0), Inches(2.55), line=WARN)
card_text(c1, [
    [("Why this is the next step", 15, True, WARN)],
    [("Our services differ — different domains, functions and structures. There is no single "
      "“correct” prompt, and a fixed evaluation checklist built for one service will not fairly "
      "judge another.", 12.5, False, INK)],
    [("So the goal is a repeatable METHOD, not a one-size template: try several prompts, then let "
      "each repository judge its own pack.", 12.5, False, INK)],
])
c2 = card(Inches(6.85), Inches(1.85), Inches(6.0), Inches(2.55), line=ACCENT)
card_text(c2, [
    [("The experiment", 15, True, ACCENT)],
    [("1  Run several prompt strategies with Claude Code to produce candidate packs for a repo", 12.5, False, INK)],
    [("2  Score each pack against that repo’s actual code (see below)", 12.5, False, INK)],
    [("3  Keep the winner for the repo, and converge on a default prompt that travels well", 12.5, False, INK)],
    [("4  Repeat on a second, different service to confirm the method — not just one result", 12.5, False, INK)],
])

# Bottom band — the codebase-derived evaluation
c3 = card(Inches(0.6), Inches(4.6), Inches(12.25), Inches(1.95), fill=ACCENT_LIGHT)
card_text(c3, [
    [("A test that fits any codebase: derive the questions from the repo itself", 15, True, ACCENT_DARK)],
    [("• Pick real functions / endpoints that exist in THAT repo and check the pack lets Duo/Copilot find and explain them",
      12.5, False, INK)],
    [("• Ask the AI to plan a change to one of those existing functions using only the pack — does it name the right files and callers?",
      12.5, False, INK)],
    [("• Score on: accuracy vs the real code · coverage of the repo’s real functions · size fits the chat context · the AI can act on it",
      12.5, False, INK)],
    [("Because the questions come from each repository’s own functions, the same method works for services that do very different things.",
      12.5, True, ACCENT_DARK)],
])

box = textbox(Inches(11.9), Inches(7.05), Inches(1.1), Inches(0.35))
set_text(box.text_frame, [[("Next", 11, False, MUTED)]], align=PP_ALIGN.RIGHT)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "next-steps-slide.pptx")
prs.save(out)
print(f"Wrote {out}")
